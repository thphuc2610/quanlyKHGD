import streamlit as st
import pandas as pd
import io
import re
from openpyxl.utils.dataframe import dataframe_to_rows

# =====================================================
# 1. CẤU HÌNH GIAO DIỆN
# =====================================================
st.set_page_config(
    page_title="Hệ thống Quản lý Giảng dạy V28.0",
    layout="wide",
    page_icon="💎"
)

st.markdown("""
<style>
thead tr th {
    background-color: #1E3A8A !important;
    color: white !important;
    font-weight: bold !important;
    text-align: center !important;
}
.main { background-color: #f8f9fa; }
[data-testid="stMetric"] {
    background-color: white;
    padding: 20px;
    border-radius: 12px;
    box-shadow: 0 4px 10px rgba(0,0,0,0.05);
    border-top: 5px solid #1E3A8A;
}
</style>
""", unsafe_allow_html=True)

# =====================================================
# 2. CHUẨN HÓA TÊN GIẢNG VIÊN
# =====================================================
def normalize_name(name):
    name = str(name)
    if name == 'nan' or not name.strip():
        return ""
    name = re.sub(r"\(.*?\)", "", name)
    name = re.sub(r"^\s*(?:(?:GS|PGS|TS|ThS|CN|KS)\.?\s*)+", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+", " ", name)
    return name.strip()

# =====================================================
# 3. TÍNH HỆ SỐ K
# =====================================================
def calculate_k(row):
    ten_mon = str(row.get('HP', '')).lower()
    sv = row.get('SV', 0)
    bm = str(row.get('BM', '')).strip()
    don_vi = str(row.get('Đơn vị', '')).strip()

    if not ten_mon:
        return 1.0

    if 'đồ án' in ten_mon:
        k = 1.1 + (sv - 40) * 0.01
        return round(max(1.0, min(1.5, k)), 2)

    if (
        bm in [
            "Bộ môn Ngôn ngữ Anh",
            "Bộ môn Tiếng anh",
            "Bộ môn Ngôn ngữ Trung Quốc"
        ]
        and don_vi == "KHCB"
    ):
        k = 1.1 + (sv - 30) * 0.01
        return round(max(0.9, min(1.5, k)), 2)

    if (
        bm == "Bộ môn Giáo dục thể chất"
        and don_vi == "KHCB"
    ):
        k = 1 + (sv - 50) * 0.01
        return round(max(0.9, min(1.2, k)), 2)

    if 'thí nghiệm' in ten_mon:
        k = 0.6 + (0.6 + (sv - 25 - 25) * 0.015)
        return round(k, 2)

    if 'thực tập trắc địa' in ten_mon:
        if sv < 20:
            return 2
        k = 2.5 + (sv - 25) * 0.05
        return round(k, 2)

    if 'thực tập tốt nghiệp' in ten_mon or 'thực tập nghề nghiệp' in ten_mon:
        return 0.5

    if 'thực tập' in ten_mon:
        if sv < 30:
            return 2
        k = 2.5 + (sv - 35) * 0.05
        return round(k, 2)

    k = 1.0 + (sv - 40) * 0.01
    return round(max(0.9, min(1.5, k)), 2)

def calculate_k_lt_th(row):
    sv = row.get('SV', 0)
    bm = str(row.get('BM', '')).strip()
    don_vi = str(row.get('Đơn vị', '')).strip()
    if bm == "Bộ môn Giáo dục thể chất" and don_vi == "KHCB":
        k_lt = 1 + (sv - 50) * 0.01
        k_lt = max(0.9, min(1.2, k_lt))
        k_th = 0.7 + (sv - 50) * 0.01
        k_th = max(0.6, min(0.9, k_th))
        return round(k_lt, 2), round(k_th, 2)
    return 0.0, 0.0

# =====================================================
# 5. MAIN
# =====================================================
def main():
    # ===== SIDEBAR + AVATAR =====
    with st.sidebar:
        if 'avatar_data' not in st.session_state:
            st.session_state.avatar_data = None

        if st.session_state.avatar_data is None:
            avatar_file = st.file_uploader("Chọn ảnh đại diện", type=["jpg", "png", "jpeg"])
            if avatar_file:
                st.session_state.avatar_data = avatar_file.read()
                st.rerun()
        else:
            st.image(st.session_state.avatar_data, width=150)
            if st.button("🔄 Thay đổi Avatar"):
                st.session_state.avatar_data = None
                st.rerun()

        st.markdown("### **TS. Phan Khanh Khánh**")
        st.caption("📧 khanhpk@tlu.edu.vn")
        st.markdown("---")

        che_do = st.radio(
            "Chế độ hiển thị:",
            ["Tổng hợp giảng dạy", "Chi tiết từng giảng viên"],
            label_visibility="collapsed"
        )

        st.markdown("---")
        uploaded_files = st.file_uploader(
            "Tải tệp Excel công việc",
            type=["xlsx"],
            accept_multiple_files=True
        )

    if not uploaded_files:
        st.title("💎 Prestige Teaching Dashboard")
        st.info("Vui lòng tải file Excel để bắt đầu.")
        return

    all_data = []
    hptn_data_list = []

    for uploaded_file in uploaded_files:
        xl = pd.ExcelFile(uploaded_file)
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            df.columns = [str(c).lower().strip() for c in df.columns]

            def find_col(keywords):
                for c in df.columns:
                    if any(k in c for k in keywords):
                        return c
                return None

            c_gv = find_col(['giảng viên'])
            c_hp = find_col(['học phần', 'tên lớp'])
            c_tc = find_col(['tín chỉ', 'tc'])
            c_sv = find_col(['sĩ số', 'sv'])
            c_dv = find_col(['đơn vị'])
            
            c_bm_ph = find_col(['tên bộ môn ph', 'tên bộ môn', 'bộ môn'])

            # HPTN check
            c_gv1 = find_col(['gv hướng dẫn 1'])
            c_gv2 = find_col(['gv hướng dẫn 2'])
            c_detai = find_col(['tên đề tài'])

            if c_detai and (c_gv1 or c_gv2):
                gv_col = c_gv1 if c_gv1 else c_gv2
                df_hptn = pd.DataFrame({
                    'GV': df[gv_col].astype(str).apply(normalize_name),
                    'Tên Đề Tài': df[c_detai]
                })
                df_hptn = df_hptn[df_hptn['GV'] != ""]
                df_hptn = df_hptn.dropna(subset=['GV', 'Tên Đề Tài'])
                hptn_data_list.append(df_hptn)
            
            elif c_gv and c_hp:
                if not c_bm_ph:
                    st.warning(f"Không tìm thấy cột 'Bộ môn' trong sheet {sheet_name}. Sẽ để trống.")
                
                temp = pd.DataFrame({
                    'GV': df[c_gv].ffill().astype(str).apply(normalize_name),
                    'HP': df[c_hp],
                    'TC': pd.to_numeric(df[c_tc], errors='coerce').fillna(0),
                    'SV': pd.to_numeric(df[c_sv], errors='coerce').fillna(0),
                    'BM': df[c_bm_ph].astype(str).str.strip() if c_bm_ph else "",
                    'Đơn vị': df[c_dv].astype(str).str.strip() if c_dv else ""
                })
                temp = temp.dropna(subset=['HP'])
                all_data.append(temp)

    # Process HPTN
    hptn_summary = pd.DataFrame(columns=['STT', 'Giảng viên', 'Đơn vị', 'Số lượng', 'Quy chuẩn'])
    if hptn_data_list:
        full_hptn = pd.concat(hptn_data_list, ignore_index=True)
        hptn_grouped = full_hptn.groupby('GV', as_index=False).size()
        hptn_summary['Giảng viên'] = hptn_grouped['GV']
        hptn_summary['Số lượng'] = hptn_grouped['size']
        hptn_summary['STT'] = range(1, len(hptn_summary) + 1)
        hptn_summary['Đơn vị'] = ""
        # The 'Quy chuẩn' will be injected via openpyxl

    # Process HĐGD
    full_data = pd.DataFrame()
    if all_data:
        full_data = pd.concat(all_data, ignore_index=True)
        full_data['Hệ số K'] = full_data.apply(calculate_k, axis=1)
        
        # Calculate K_lt and K_th
        full_data[['Hệ số K_lt', 'Hệ số K_th']] = full_data.apply(
            lambda row: pd.Series(calculate_k_lt_th(row)), axis=1
        )
        # We don't need Tiết quy đổi column if the formula computes GC directly, but we can keep it if needed.
        
    st.header("📊 Khối lượng giảng dạy")

    # Split into CHI TIẾT_KLGD and CHI TIẾT_KLGD CÁC MÔN GDTC
    df_ct_gdtc = pd.DataFrame()
    df_ct_normal = pd.DataFrame()
    
    if not full_data.empty:
        mask_gdtc = full_data['BM'] == "Bộ môn Giáo dục thể chất"
        df_ct_gdtc = full_data[mask_gdtc].copy()
        df_ct_normal = full_data.copy()
        
        # Format df_ct_normal
        df_ct_normal_out = pd.DataFrame({
            'STT': range(1, len(df_ct_normal) + 1),
            'Giảng viên': df_ct_normal['GV'],
            'Tên lớp': df_ct_normal['HP'],
            'TC': df_ct_normal['TC'],
            'SV đăng ký': df_ct_normal['SV'],
            'Tên bộ môn PH': df_ct_normal['BM'],
            'Đơn vị': df_ct_normal['Đơn vị'],
            'Hệ số K': df_ct_normal['Hệ số K'],
            'Hệ số K_lt': df_ct_normal['Hệ số K_lt'],
            'Hệ số K_th': df_ct_normal['Hệ số K_th'],
            'GC': 0.0, 
            'Hệ số đúng': "" 
        })
        # Calculate GC for normal
        def compute_gc(row):
            if row['Tên bộ môn PH'] == "Bộ môn Giáo dục thể chất" and row['Đơn vị'] == "KHCB":
                return row['Hệ số K_lt'] * 10 + row['Hệ số K_th'] * 20
            else:
                return row['TC'] * 15 * row['Hệ số K']
        df_ct_normal_out['GC'] = df_ct_normal_out.apply(compute_gc, axis=1)

        # Format df_ct_gdtc
        df_ct_gdtc_out = pd.DataFrame({
            'STT': range(1, len(df_ct_gdtc) + 1),
            'Giảng viên': df_ct_gdtc['GV'],
            'Tên lớp': df_ct_gdtc['HP'],
            'TC': df_ct_gdtc['TC'],
            'SV đăng ký': df_ct_gdtc['SV'],
            'Tên bộ môn PH': df_ct_gdtc['BM'],
            'Đơn vị': df_ct_gdtc['Đơn vị'],
            'Hệ số K': df_ct_gdtc['Hệ số K'],
            'Hệ số K_lt': df_ct_gdtc['Hệ số K_lt'],
            'Hệ số K_th': df_ct_gdtc['Hệ số K_th'],
            'GC': "" # Formula in Excel
        })
    else:
        df_ct_normal_out = pd.DataFrame()
        df_ct_gdtc_out = pd.DataFrame()

    # Create Summary dataframe for display and base Excel
    # We gather all unique teachers
    all_gvs = set()
    if not df_ct_normal_out.empty:
        all_gvs.update(df_ct_normal_out['Giảng viên'].unique())
    if not df_ct_gdtc_out.empty:
        all_gvs.update(df_ct_gdtc_out['Giảng viên'].unique())
    if not hptn_summary.empty:
        all_gvs.update(hptn_summary['Giảng viên'].unique())
        
    all_gvs = sorted(list(all_gvs))
    df_tonghop = pd.DataFrame({
        'STT': range(1, len(all_gvs) + 1),
        'Giảng viên': all_gvs,
        'Số lớp': 0,
        'Số TC': 0,
        'Tổng SV': 0,
        'Quy chuẩn HĐGD': 0,
        'Khối lượng HPTN': 0,
        'Tổng cộng': 0
    })

    if che_do == "Tổng hợp giảng dạy":
        st.dataframe(df_tonghop, use_container_width=True)
    else:
        st.dataframe(full_data, use_container_width=True)

    # ===== EXPORT EXCEL WITH OPENPYXL =====
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df_tonghop.to_excel(writer, sheet_name='TỔNG HỢP_KLGD', index=False)
        if not df_ct_normal_out.empty:
            df_ct_normal_out.to_excel(writer, sheet_name='CHI TIẾT_KLGD', index=False)
        if not df_ct_gdtc_out.empty:
            df_ct_gdtc_out.to_excel(writer, sheet_name='CHI TIẾT_KLGD CÁC MÔN GDTC', index=False)
        if not hptn_summary.empty:
            hptn_summary.to_excel(writer, sheet_name='KHỐI LƯỢNG HPTN', index=False)
            
        workbook = writer.book
        
        # Inject formulas into KHỐI LƯỢNG HPTN
        if 'KHỐI LƯỢNG HPTN' in workbook.sheetnames:
            ws_hptn = workbook['KHỐI LƯỢNG HPTN']
            for row_idx in range(2, ws_hptn.max_row + 1):
                # E = D * 14
                ws_hptn[f'E{row_idx}'] = f'=D{row_idx}*14'

        # Inject formulas into CHI TIẾT_KLGD CÁC MÔN GDTC
        if 'CHI TIẾT_KLGD CÁC MÔN GDTC' in workbook.sheetnames:
            ws_gdtc = workbook['CHI TIẾT_KLGD CÁC MÔN GDTC']
            for row_idx in range(2, ws_gdtc.max_row + 1):
                # K = I * 10 + J * 20
                ws_gdtc[f'K{row_idx}'] = f'=I{row_idx}*10+J{row_idx}*20'
                
        # Inject formulas into TỔNG HỢP_KLGD
        ws_tonghop = workbook['TỔNG HỢP_KLGD']
        for row_idx in range(2, ws_tonghop.max_row + 1):
            row_str = str(row_idx)
            # C: Số lớp
            ws_tonghop[f'C{row_str}'] = f"=COUNTIF('CHI TIẾT_KLGD'!B:B, B{row_str})"
            # D: Số TC
            ws_tonghop[f'D{row_str}'] = f"=SUMIF('CHI TIẾT_KLGD'!B:B, B{row_str}, 'CHI TIẾT_KLGD'!D:D)"
            # E: Tổng SV
            ws_tonghop[f'E{row_str}'] = f"=SUMIF('CHI TIẾT_KLGD'!B:B, B{row_str}, 'CHI TIẾT_KLGD'!E:E)"
            # F: Quy chuẩn HĐGD
            ws_tonghop[f'F{row_str}'] = f"=SUMIF('CHI TIẾT_KLGD'!B:B, B{row_str}, 'CHI TIẾT_KLGD'!K:K)"
            # G: Khối lượng HPTN
            ws_tonghop[f'G{row_str}'] = f"=SUMIF('KHỐI LƯỢNG HPTN'!B:B, B{row_str}, 'KHỐI LƯỢNG HPTN'!E:E)"
            # H: Tổng cộng
            ws_tonghop[f'H{row_str}'] = f"=F{row_str}+G{row_str}"

    st.sidebar.download_button(
        "📥 Tải báo cáo Excel",
        data=buffer.getvalue(),
        file_name="bao-cao-khoi-luong-giang-day.xlsx"
    )

if __name__ == "__main__":
    main()
